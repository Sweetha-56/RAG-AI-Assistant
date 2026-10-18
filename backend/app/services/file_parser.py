import os
import re
import io
import base64
import logging
from typing import List, Dict, Any, Optional

import requests
import fitz
import pdfplumber
from PIL import Image
from pypdf import PdfReader
from docx import Document
import pandas as pd

from app.core.config import settings
from app.services.ocr_service import extract_text_from_image, extract_text_from_pil_image

logger = logging.getLogger("file_parser")

try:
    from pptx import Presentation
except Exception:
    Presentation = None


MEDICINE_NAMES = [
    "Tagrisso", "Imfinzi", "Calquence", "Lynparza", "Enhertu",
    "Zoladex", "Truqap", "Imjudo", "Datroway", "Farxiga",
    "Crestor", "Brilinta", "Seloken", "Lokelma", "roxadustat",
    "Wainua", "Symbicort", "Fasenra", "Breztri", "Tezspire",
    "Beyfortus", "Synagis", "FluMist", "Ultomiris", "Soliris",
    "Strensiq", "Koselugo", "Nexium", "Saphnelo", "Airsupra",
    "Pulmicort", "Wainzua", "Beyonttra", "Voydeya"
]

KNOWN_SECTION_HEADINGS = [
    "Summary",
    "Revenue Drivers",
    "R&D Progress",
    "Sustainability",
    "Financial Performance",
    "Financial Statements",
    "Glossary",
    "Results highlights",
    "Operating and financial review",
    "Other shareholder information",
    "Notes to the Interim financial statements",
    "Rare Disease",
    "Oncology",
    "BioPharmaceuticals",
    "CVRM",
    "R&I",
    "V&I",
    "Interim financial statements",
    "Financial calendar",
]


def clean_extracted_text(text: str) -> str:
    if not text:
        return ""

    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[^\S\n]+", " ", text)
    text = re.sub(r"[•●▪◦]", "-", text)
    return text.strip()


def extract_known_headings_from_text(text: str) -> List[str]:
    found = []
    normalized_text = clean_extracted_text(text)

    for heading in KNOWN_SECTION_HEADINGS:
        if heading.lower() in normalized_text.lower():
            found.append(heading)

    return found


def extract_table_titles_from_text(text: str) -> List[str]:
    titles = []
    lines = [line.strip() for line in text.split("\n") if line.strip()]

    for line in lines:
        if re.match(r"^Table\s+\d+[:.]?", line, flags=re.IGNORECASE):
            titles.append(line)

    return titles


def is_useless_image_description(text: str) -> bool:
    if not text:
        return True

    normalized = text.strip().lower()

    bad_patterns = [
        "i'm unable to view images directly",
        "i am unable to view images directly",
        "please provide a description of the image",
        "i cannot view images directly",
        "as a text-based model",
        "overall style",
    ]

    return any(p in normalized for p in bad_patterns)


def call_llava_for_image_bytes(image_bytes: bytes, prompt: str) -> str:
    if not settings.enable_image_description:
        return ""

    if not image_bytes:
        return ""

    if hasattr(settings, "llava_max_request_image_bytes"):
        if len(image_bytes) > settings.llava_max_request_image_bytes:
            logger.warning(
                f"Skipping LLaVA call because image bytes ({len(image_bytes)}) exceed "
                f"configured safe limit ({settings.llava_max_request_image_bytes})"
            )
            return ""

    try:
        image_b64 = base64.b64encode(image_bytes).decode("utf-8")

        response = requests.post(
            f"{settings.ollama_base_url}/api/generate",
            json={
                "model": settings.ollama_vision_model,
                "prompt": prompt,
                "images": [image_b64],
                "stream": False,
                "options": {
                    "temperature": 0
                }
            },
            timeout=settings.llm_timeout
        )
        response.raise_for_status()
        data = response.json()

        text = clean_extracted_text(data.get("response", "") or "")
        if is_useless_image_description(text):
            return ""

        return text

    except Exception as e:
        logger.warning(f"LLaVA image description failed: {e}")
        return ""


def call_llava_for_image(file_path: str) -> str:
    try:
        with open(file_path, "rb") as f:
            image_bytes = f.read()

        prompt = (
            "Read this image carefully. Extract the important visible text if present. "
            "Identify the main topic, labels, concepts, and any key numeric or factual information. "
            "Return a concise factual description useful for semantic retrieval. "
            "Do not say you cannot view the image."
        )
        return call_llava_for_image_bytes(image_bytes, prompt)
    except Exception as e:
        logger.warning(f"LLaVA failed for file {file_path}: {e}")
        return ""


def build_structured_image_text(ocr_text: str, image_description: str) -> str:
    parts = []

    if ocr_text:
        parts.append(f"OCR Text:\n{ocr_text}")

    if image_description:
        parts.append(f"Image Summary:\n{image_description}")

    return "\n\n".join(parts).strip()


def try_numeric_summary(df: pd.DataFrame, sheet_name: str) -> str:
    summary_lines = []
    numeric_cols = []

    for col in df.columns:
        series = pd.to_numeric(df[col], errors="coerce")
        if series.notna().sum() > 0:
            numeric_cols.append((col, series))

    if not numeric_cols:
        return ""

    summary_lines.append(f"Numeric Summary for Sheet: {sheet_name}")

    for col, series in numeric_cols:
        valid = series.dropna()
        if valid.empty:
            continue

        max_idx = valid.idxmax()
        min_idx = valid.idxmin()

        summary_lines.append(
            f"Column '{col}': count={int(valid.count())}, "
            f"min={valid.min()}, max={valid.max()}, avg={round(valid.mean(), 2)}"
        )

        try:
            row_max = df.loc[max_idx].to_dict()
            summary_lines.append(f"Highest row for '{col}': {row_max}")
        except Exception:
            pass

        try:
            row_min = df.loc[min_idx].to_dict()
            summary_lines.append(f"Lowest row for '{col}': {row_min}")
        except Exception:
            pass

    return clean_extracted_text("\n".join(summary_lines))


def dataframe_to_retrieval_text(df: pd.DataFrame, sheet_name: str) -> str:
    df = df.fillna("")

    lines = [f"Sheet Name: {sheet_name}"]

    columns = [str(c).strip() for c in df.columns.tolist()]
    if columns:
        lines.append("Columns: " + ", ".join(columns))

    max_rows = min(len(df), 200)
    for idx, row in df.head(max_rows).iterrows():
        row_parts = []
        for col in df.columns:
            col_name = str(col).strip()
            value = str(row[col]).strip()

            if value and value.lower() != "nan":
                row_parts.append(f"{col_name}: {value}")

        if row_parts:
            lines.append(f"Row {idx + 1}: " + " | ".join(row_parts))

    numeric_summary = try_numeric_summary(df, sheet_name)
    if numeric_summary:
        lines.append(numeric_summary)

    return clean_extracted_text("\n".join(lines))


def render_pdf_page_to_pil(page, dpi: int = 150) -> Image.Image:
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix, alpha=False)
    image = Image.open(io.BytesIO(pix.tobytes("png")))
    return image


def compress_image_for_llava(image: Image.Image, max_width: int = 1200, jpeg_quality: int = 70) -> bytes:
    img = image.convert("RGB")

    if img.width > max_width:
        ratio = max_width / float(img.width)
        new_height = int(img.height * ratio)
        img = img.resize((max_width, new_height))

    buffer = io.BytesIO()
    img.save(buffer, format="JPEG", quality=jpeg_quality, optimize=True)
    return buffer.getvalue()


def build_pdf_page_vision_prompt(page_number: int) -> str:
    return (
        f"Read this document page carefully. "
        f"Extract important visible text, figure titles, chart titles, axis labels, table headings, "
        f"key values, percentages, and the main takeaway. "
        f"If the page contains a chart or graph, identify the chart type and describe the trend. "
        f"If there is a table, describe its structure and important figures. "
        f"Return concise factual retrieval text for page {page_number}. "
        f"Do not say you cannot view the image."
    )


def extract_pdf_tables(file_path: str) -> List[Dict[str, Any]]:
    entries = []

    if not settings.enable_table_extraction:
        return entries

    try:
        with pdfplumber.open(file_path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                try:
                    tables = page.extract_tables() or []
                except Exception:
                    tables = []

                page_text = ""
                try:
                    page_text = clean_extracted_text(page.extract_text() or "")
                except Exception:
                    page_text = ""

                section_headings = extract_known_headings_from_text(page_text)
                table_titles = extract_table_titles_from_text(page_text)
                table_title = table_titles[0] if table_titles else ""

                for table_index, table in enumerate(tables, start=1):
                    if not table:
                        continue

                    rows = []
                    for row in table:
                        row = row or []
                        cells = [clean_extracted_text(str(cell or "")) for cell in row]
                        if any(cells):
                            rows.append(" | ".join(cells))

                    table_text = "\n".join(rows).strip()
                    if table_text:
                        entries.append({
                            "type": "pdf_table",
                            "content_type": "table",
                            "page_number": page_index,
                            "table_number": table_index,
                            "table_title": table_title,
                            "table_titles": table_titles,
                            "section_headings": section_headings,
                            "text": table_text,
                            "image_paths": []
                        })
    except Exception as e:
        logger.warning(f"PDF table extraction failed for {file_path}: {e}")

    return entries


def normalize_medicine_block(medicine_name: str, raw_text: str) -> str:
    text = clean_extracted_text(raw_text)
    lines = [line.strip() for line in text.split("\n") if line.strip()]

    normalized_lines = [f"Medicine: {medicine_name}"]

    for line in lines:
        if line.lower() == medicine_name.lower():
            continue
        normalized_lines.append(line)

    joined = "\n".join(normalized_lines)
    joined = joined.replace("Q1 2025", "\nPeriod: Q1 2025\n")
    joined = joined.replace("Total Revenue", "\nMetric: Total Revenue\n")
    joined = joined.replace("% Change", "\nPercent Change Section\n")

    return clean_extracted_text(joined)


def _extract_number_after_label(text: str, label: str) -> Optional[str]:
    pattern = rf"{re.escape(label)}\s+([\d,]+|\([\d]+\)|>?\d+x|n/m|-)"
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return match.group(1) if match else None


def build_medicine_summary(medicine_name: str, raw_text: str) -> str:
    text = clean_extracted_text(raw_text)

    total_value = _extract_number_after_label(text, "Total")
    us_value = _extract_number_after_label(text, "US")
    europe_value = _extract_number_after_label(text, "Europe")
    em_value = _extract_number_after_label(text, "Emerging Markets")
    row_value = _extract_number_after_label(text, "Established RoW")

    actual_matches = re.findall(
        r"Total\s+[\d,]+\s+([()>\dxmn/\-]+)\s+([()>\dxmn/\-]+)",
        text,
        flags=re.IGNORECASE
    )
    actual_change = actual_matches[0][0] if actual_matches else ""
    cer_change = actual_matches[0][1] if actual_matches else ""

    summary_lines = [
        f"Medicine Summary: {medicine_name}",
        "Period: Q1 2025"
    ]

    if total_value:
        summary_lines.append(f"Total revenue was {total_value}.")

    region_parts = []
    if us_value:
        region_parts.append(f"US {us_value}")
    if europe_value:
        region_parts.append(f"Europe {europe_value}")
    if em_value:
        region_parts.append(f"Emerging Markets {em_value}")
    if row_value:
        region_parts.append(f"Established RoW {row_value}")

    if region_parts:
        summary_lines.append("Regional contribution included: " + ", ".join(region_parts) + ".")

    if actual_change or cer_change:
        summary_lines.append(
            f"Reported growth/decline was Actual {actual_change or 'N/A'} and CER {cer_change or 'N/A'}."
        )

    commentary_lines = []
    for line in text.split("\n"):
        line = line.strip()
        if (
            line.startswith("-")
            or "growth" in line.lower()
            or "decline" in line.lower()
            or "competition" in line.lower()
            or "demand" in line.lower()
            or "conversion" in line.lower()
        ):
            commentary_lines.append(line)

    if commentary_lines:
        summary_lines.append("Key commentary: " + " ".join(commentary_lines[:3]))

    return clean_extracted_text("\n".join(summary_lines))


def extract_medicine_blocks_from_text(page_text: str, page_number: int):
    entries = []
    if not page_text:
        return entries

    pattern = r"(?=\n?(?:" + "|".join(re.escape(m) for m in MEDICINE_NAMES) + r")\s*\n)"
    parts = re.split(pattern, "\n" + page_text)

    section_headings = extract_known_headings_from_text(page_text)
    table_titles = extract_table_titles_from_text(page_text)

    for part in parts:
        part = clean_extracted_text(part)
        if not part:
            continue

        matched_name = None
        for med in MEDICINE_NAMES:
            if part.lower().startswith(med.lower()):
                matched_name = med
                break

        if matched_name:
            normalized = normalize_medicine_block(matched_name, part)
            entries.append({
                "type": "pdf_medicine_block",
                "content_type": "table",
                "page_number": page_number,
                "text": normalized,
                "medicine_name": matched_name,
                "section_headings": section_headings,
                "table_titles": table_titles,
                "image_paths": []
            })

            summary_text = build_medicine_summary(matched_name, part)
            if summary_text:
                entries.append({
                    "type": "pdf_medicine_summary",
                    "content_type": "summary",
                    "page_number": page_number,
                    "text": summary_text,
                    "medicine_name": matched_name,
                    "section_headings": section_headings,
                    "table_titles": table_titles,
                    "image_paths": []
                })

    return entries


def parse_pdf(file_path: str):
    entries = []

    try:
        reader = PdfReader(file_path)
    except Exception as e:
        logger.warning(f"Failed to open PDF with pypdf {file_path}: {e}")
        reader = None

    native_text_by_page = {}
    if reader is not None:
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = clean_extracted_text(page.extract_text() or "")
            except Exception as e:
                logger.warning(f"Failed extracting native text from page {i} in {file_path}: {e}")
                text = ""

            native_text_by_page[i] = text

            if text:
                section_headings = extract_known_headings_from_text(text)
                table_titles = extract_table_titles_from_text(text)

                entries.append({
                    "type": "pdf_text",
                    "content_type": "text",
                    "page_number": i,
                    "text": text,
                    "section_headings": section_headings,
                    "table_titles": table_titles,
                    "image_paths": []
                })

                medicine_entries = extract_medicine_blocks_from_text(text, i)
                entries.extend(medicine_entries)

    entries.extend(extract_pdf_tables(file_path))

    try:
        pdf_doc = fitz.open(file_path)
    except Exception as e:
        logger.warning(f"Failed to open PDF with PyMuPDF {file_path}: {e}")
        return entries

    os.makedirs(settings.temp_image_dir, exist_ok=True)

    total_pages = len(pdf_doc)
    max_pages = min(total_pages, settings.max_pdf_pages_for_vision)

    logger.info(f"PDF {os.path.basename(file_path)} total_pages={total_pages}, vision_limit={max_pages}")

    for page_index in range(total_pages):
        page_number = page_index + 1
        page = pdf_doc.load_page(page_index)

        native_text = native_text_by_page.get(page_number, "")
        section_headings = extract_known_headings_from_text(native_text)
        table_titles = extract_table_titles_from_text(native_text)

        ocr_text = ""
        visual_summary = ""
        saved_page_image_path = ""

        should_render_page = False
        should_do_ocr = False
        should_do_vision = False

        if settings.enable_pdf_page_ocr and len(native_text) < settings.min_pdf_text_chars_before_ocr:
            should_render_page = True
            should_do_ocr = True

        if (
            settings.enable_pdf_page_image_summary
            and page_index < max_pages
            and len(native_text) < 500
        ):
            should_render_page = True
            should_do_vision = True

        page_image = None
        if should_render_page:
            try:
                page_image = render_pdf_page_to_pil(page, dpi=settings.pdf_render_dpi)

                saved_page_image_path = os.path.join(
                    settings.temp_image_dir,
                    f"{os.path.splitext(os.path.basename(file_path))[0]}_page_{page_number}.png"
                )

                try:
                    page_image.save(saved_page_image_path)
                except Exception:
                    saved_page_image_path = ""
            except Exception as e:
                logger.warning(f"Failed rendering PDF page {page_number} in {file_path}: {e}")
                page_image = None

        if should_do_ocr and page_image is not None:
            try:
                ocr_text = clean_extracted_text(extract_text_from_pil_image(page_image, psm=6))
                if ocr_text:
                    entries.append({
                        "type": "pdf_ocr",
                        "content_type": "ocr",
                        "page_number": page_number,
                        "text": ocr_text,
                        "ocr_text": ocr_text,
                        "section_headings": section_headings,
                        "table_titles": table_titles,
                        "image_paths": [saved_page_image_path] if saved_page_image_path else []
                    })
            except Exception as e:
                logger.warning(f"OCR failed for PDF page {page_number} in {file_path}: {e}")

        if should_do_vision and page_image is not None:
            try:
                max_width = getattr(settings, "llava_max_image_width", 1200)
                jpeg_quality = getattr(settings, "llava_jpeg_quality", 70)

                page_image_bytes = compress_image_for_llava(
                    page_image,
                    max_width=max_width,
                    jpeg_quality=jpeg_quality
                )

                prompt = build_pdf_page_vision_prompt(page_number)
                visual_summary = clean_extracted_text(
                    call_llava_for_image_bytes(page_image_bytes, prompt)
                )

                if visual_summary:
                    entries.append({
                        "type": "pdf_visual_summary",
                        "content_type": "chart_or_visual",
                        "page_number": page_number,
                        "text": visual_summary,
                        "image_description": visual_summary,
                        "section_headings": section_headings,
                        "table_titles": table_titles,
                        "image_paths": [saved_page_image_path] if saved_page_image_path else []
                    })
            except Exception as e:
                logger.warning(f"Vision summary failed for PDF page {page_number} in {file_path}: {e}")

        hybrid_parts = []
        if native_text:
            hybrid_parts.append(f"Native Text:\n{native_text[:2500]}")
        if ocr_text:
            hybrid_parts.append(f"OCR Text:\n{ocr_text[:2000]}")
        if visual_summary:
            hybrid_parts.append(f"Visual Summary:\n{visual_summary}")

        hybrid_text = clean_extracted_text("\n\n".join(hybrid_parts))
        if hybrid_text:
            entries.append({
                "type": "pdf_page_summary",
                "content_type": "page_summary",
                "page_number": page_number,
                "text": hybrid_text,
                "ocr_text": ocr_text,
                "image_description": visual_summary,
                "section_headings": section_headings,
                "table_titles": table_titles,
                "image_paths": [saved_page_image_path] if saved_page_image_path else []
            })

    pdf_doc.close()
    logger.info(f"Finished parse_pdf for {os.path.basename(file_path)} with total entries={len(entries)}")
    return entries


def parse_docx(file_path: str):
    doc = Document(file_path)
    entries = []

    para_buffer = []
    for p in doc.paragraphs:
        txt = clean_extracted_text(p.text or "")
        if txt:
            para_buffer.append(txt)

    if para_buffer:
        text = "\n\n".join(para_buffer)
        entries.append({
            "type": "docx_text",
            "content_type": "text",
            "page_number": None,
            "text": text,
            "section_headings": extract_known_headings_from_text(text),
            "table_titles": extract_table_titles_from_text(text),
            "image_paths": []
        })

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [clean_extracted_text(cell.text or "") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))

        table_text = "\n".join(rows).strip()
        if table_text:
            entries.append({
                "type": "docx_table",
                "content_type": "table",
                "page_number": None,
                "table_number": table_index,
                "text": table_text,
                "section_headings": extract_known_headings_from_text(table_text),
                "table_titles": extract_table_titles_from_text(table_text),
                "image_paths": []
            })

    return entries


def parse_xlsx(file_path: str):
    entries = []

    try:
        xls = pd.ExcelFile(file_path)
    except Exception as e:
        logger.warning(f"Failed to open Excel file {file_path}: {e}")
        return entries

    for sheet_name in xls.sheet_names:
        try:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
        except Exception as e:
            logger.warning(f"Failed to read sheet '{sheet_name}' in {file_path}: {e}")
            continue

        if df.empty and len(df.columns) == 0:
            continue

        text = dataframe_to_retrieval_text(df, sheet_name)

        if text:
            entries.append({
                "type": "xlsx_sheet",
                "content_type": "table",
                "sheet_name": sheet_name,
                "page_number": None,
                "text": text,
                "section_headings": [],
                "table_titles": [],
                "image_paths": []
            })

    return entries


def parse_txt(file_path: str):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = clean_extracted_text(f.read())

    return [{
        "type": "txt",
        "content_type": "text",
        "page_number": None,
        "text": text,
        "section_headings": extract_known_headings_from_text(text),
        "table_titles": extract_table_titles_from_text(text),
        "image_paths": []
    }] if text else []


def parse_image(file_path: str):
    ocr_text = clean_extracted_text(extract_text_from_image(file_path))
    image_description = clean_extracted_text(call_llava_for_image(file_path))
    combined_text = build_structured_image_text(ocr_text, image_description)

    return [{
        "type": "image_content",
        "content_type": "image",
        "page_number": None,
        "text": combined_text,
        "ocr_text": ocr_text,
        "image_description": image_description,
        "section_headings": [],
        "table_titles": [],
        "image_paths": [file_path]
    }] if combined_text else []


def parse_pptx(file_path: str):
    if Presentation is None:
        raise ValueError("python-pptx is not installed")

    prs = Presentation(file_path)
    entries = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                txt = clean_extracted_text(shape.text)
                if txt:
                    parts.append(txt)

        slide_text = "\n\n".join(parts).strip()
        if slide_text:
            entries.append({
                "type": "pptx_slide",
                "content_type": "text",
                "page_number": slide_num,
                "slide_number": slide_num,
                "text": slide_text,
                "section_headings": extract_known_headings_from_text(slide_text),
                "table_titles": extract_table_titles_from_text(slide_text),
                "image_paths": []
            })

    return entries


def parse_file(file_path: str):
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return parse_pdf(file_path)
    if ext == ".docx":
        return parse_docx(file_path)
    if ext == ".xlsx":
        return parse_xlsx(file_path)
    if ext == ".txt":
        return parse_txt(file_path)
    if ext == ".pptx":
        return parse_pptx(file_path)
    if ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif"]:
        return parse_image(file_path)

    raise ValueError(f"Unsupported file type: {ext}")